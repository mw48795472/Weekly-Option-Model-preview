"""Export the reviewed Markdown to Word; requires the optional reports extra."""
import argparse
import re
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/hypotheses-review.md"
FONT = "Malgun Gothic"
TOKENS = re.compile(r"(\[[^\]]+\]\(<[^>]+>\)|\[[^\]]+\]\([^\s)]+\)|\*\*[^*]+\*\*|`[^`]+`)")


def font(run, size=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)


def resolve_link(target, ref):
    if target.startswith(("https://", "http://")):
        return target
    path = (SOURCE.parent / target).resolve()
    if not path.is_file() or not path.is_relative_to(ROOT):
        raise ValueError(f"Invalid local source link: {target}")
    relative = path.relative_to(ROOT).as_posix()
    return "https://github.com/mw48795472/Weekly-Option-Model-preview/blob/" + quote(ref, safe="") + "/" + quote(relative, safe="/")


def add_inline(paragraph, text, ref, size=None, header=False):
    for part in TOKENS.split(text):
        if not part:
            continue
        match = re.fullmatch(r"\[([^\]]+)\]\(<?(.*?)>?\)", part)
        if match:
            label, target = match.groups()
            target = target.rstrip(">")
            rel = paragraph.part.relate_to(resolve_link(target, ref), RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), rel)
            run = paragraph.add_run(label)
            font(run, size)
            run.font.color.rgb = RGBColor.from_string("185A8A")
            run.font.underline = True
            paragraph._p.remove(run._r)
            link.append(run._r)
            paragraph._p.append(link)
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        value = part[2:-2] if bold else part[1:-1] if code else part
        run = paragraph.add_run(value)
        font(run, size)
        run.bold = bold or header
        if header:
            run.font.color.rgb = RGBColor.from_string("FFFFFF")


def set_style(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)


def add_table(doc, lines, ref):
    records = [[x.strip() for x in line.strip().strip("|").split("|")] for line in lines]
    records = [r for r in records if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    count = len(records[0])
    widths = [0.5, 1.6, 2.35, 2.65] if count == 4 else [1.25, 2.75, 3.1]
    assert len(widths) == count and all(len(r) == count for r in records)
    table = doc.add_table(rows=0, cols=count)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    props = table._tbl.tblPr
    width = props.find(qn("w:tblW"))
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(round(sum(widths) * 1440)))
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn("w:" + key), val)
        borders.append(border)
    props.append(borders)
    for column, inches in zip(table.columns, widths):
        column.width = Inches(inches)
    for index, record in enumerate(records):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for col, (cell, value, inches) in enumerate(zip(row.cells, record, widths)):
            cell.width = Inches(inches)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:val"), "clear")
            shade.set(qn("w:fill"), "243B53" if index == 0 else "F2F5F8" if index % 2 == 0 else "FFFFFF")
            tcpr.append(shade)
            margin = OxmlElement("w:tcMar")
            for side, val in (("top", 90), ("bottom", 90), ("left", 100), ("right", 100)):
                tag = OxmlElement("w:" + side)
                tag.set(qn("w:w"), str(val))
                tag.set(qn("w:type"), "dxa")
                margin.append(tag)
            tcpr.append(margin)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if count == 4 and col == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, ref, size=9.5, header=index == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-ref", default="main", help="Git commit holding report and supporting evidence")
    parser.add_argument("--output", type=Path, default=ROOT / "docs/downloads/weekly_option_hypotheses_review.docx")
    args = parser.parse_args()
    doc = Document()
    # Some distributed default templates contain a blue Title border. Remove
    # paragraph borders from all styles so inherited formatting cannot restore it.
    for border in doc.styles.element.findall(".//" + qn("w:pBdr")):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.7)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.footer_distance = Inches(0.3)
    for name, size, bold in (("Normal", 11, False), ("Title", 23, True), ("Heading 1", 15, True), ("Heading 2", 12.5, True)):
        set_style(doc.styles[name], size, bold)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(7)
    normal.line_spacing = 1.3
    normal.widow_control = True
    for name in ("Heading 1", "Heading 2"):
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after = Pt(16), Pt(7)
        pf.keep_with_next = True
    doc.core_properties.title = "위클리옵션 가설 심화 검토"
    doc.core_properties.subject = "H1부터 H9 및 H9b 연구 설계와 저장자료 진단"
    doc.core_properties.author = "Weekly Option Research"
    doc.core_properties.keywords = "NH Futures, weekly options, hypothesis review"
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    font(run, 9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines, args.source_ref)
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading = line[level:].strip()
            heading = re.sub(r"[.:·]", " ", heading)
            heading = re.sub(r"\s+", " ", heading).strip()
            p = doc.add_paragraph(style="Title" if level == 1 else f"Heading {level - 1}")
            add_inline(p, heading, args.source_ref)
        else:
            p = doc.add_paragraph()
            add_inline(p, line, args.source_ref)
        index += 1
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
